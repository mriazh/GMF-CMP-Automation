"""CLI entry point for CMP Automation."""

import argparse
import asyncio
import logging
import sys
import time
from datetime import date, datetime
from pathlib import Path

from .config import Config, load_config, validate_paths
from .exceptions import CMPAutomationError, ConfigurationError
from .logging import log_run_boundary, setup_logging
from .notifier import (
    EVENT_FAILED,
    EVENT_START,
    EVENT_SUCCESS,
    Notifier,
    build_notifier,
    safe_error_category,
)
from .workflow import run_workflow

logger = logging.getLogger(__name__)

# Only these modes notify; ``generate`` stays unnotified by design.
NOTIFIED_MODES = frozenset({"full", "scrape"})


def parse_date_arg(val: str | None) -> date | None:
    """Parse YYYY-MM-DD date argument."""
    if not val or not val.strip():
        return None
    try:
        return datetime.strptime(val.strip(), "%Y-%m-%d").date()
    except ValueError as e:
        raise ConfigurationError(f"Invalid date format '{val}'. Expected YYYY-MM-DD.") from e


def parse_args() -> argparse.Namespace:
    """Parse command line arguments."""
    parser = argparse.ArgumentParser(
        description="Telkomsel CMP Portal Automation Tool - Daily Usage Query & Reporting",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "--mode",
        choices=["full", "scrape", "generate"],
        default="full",
        help="Pipeline stage: full (scrape + generate), scrape (raw/image only), or generate (existing raw/image)",
    )
    parser.add_argument(
        "--menu",
        action="store_true",
        help="Launch interactive terminal menu (Scrape, Generate, Full)",
    )
    parser.add_argument(
        "--raw-xlsx",
        type=Path,
        help="Raw portal XLSX input for --mode generate",
    )
    parser.add_argument(
        "--image",
        type=Path,
        help="Dashboard image input for --mode generate",
    )
    parser.add_argument(
        "--date",
        type=str,
        help="Target date to export and report (YYYY-MM-DD). Defaults to today in configured timezone.",
    )
    parser.add_argument(
        "--start-date",
        type=str,
        help="Start date for date range processing (YYYY-MM-DD).",
    )
    parser.add_argument(
        "--end-date",
        type=str,
        help="End date for date range processing (YYYY-MM-DD).",
    )
    parser.add_argument(
        "--headed",
        action="store_true",
        help="Run browser in headed mode (visible)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate configuration and browser launch without full workflow",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        help="Override OTP timeout in seconds",
    )
    parser.add_argument(
        "--download-dir",
        type=Path,
        help="Override download directory",
    )
    parser.add_argument(
        "--profile-dir",
        type=Path,
        help="Override Firefox profile directory",
    )
    parser.add_argument(
        "--excel-output-dir",
        dest="excel_output_dir",
        type=Path,
        help="Override monthly Excel output directory",
    )
    parser.add_argument(
        "--report-dir",
        dest="report_dir",
        type=Path,
        help="Override monthly Excel output directory",
    )
    parser.add_argument(
        "--xlsx-dir",
        dest="xlsx_dir",
        type=Path,
        help="Compatibility alias for the raw download directory",
    )
    parser.add_argument(
        "--image-dir",
        dest="image_dir",
        type=Path,
        help="Compatibility alias for dashboard image directory",
    )
    parser.add_argument(
        "--excel-template",
        type=Path,
        default=None,
        help="Override source Excel template path",
    )
    parser.add_argument(
        "--log-level",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
        default="INFO",
        help="Set logging level",
    )
    parser.add_argument(
        "--diagnose-export",
        action="store_true",
        help="Opt-in: record sanitized network metadata around export (diagnostic only)",
    )
    parser.add_argument(
        "--diagnose-auth",
        action="store_true",
        help="Opt-in: record sanitized metadata around authentication (diagnostic only)",
    )
    parser.add_argument(
        "--monitor-cycles",
        type=int,
        default=None,
        help="Optional finite dashboard refresh cycles after the first capture",
    )
    parser.add_argument(
        "--enable-connectivity",
        action="store_true",
        help="Opt in to Check Point/IMAP/WARP connectivity checks",
    )
    parser.add_argument(
        "--allow-connectivity-mutation",
        action="store_true",
        help="Allow this run to initiate or disconnect owned VPN tunnels",
    )
    return parser.parse_args()


def apply_overrides(config: Config, args: argparse.Namespace) -> None:
    """Apply CLI overrides and revalidate Pydantic constraints."""
    if args.timeout is not None:
        if not 10 <= args.timeout <= 600:
            raise ConfigurationError("--timeout must be between 10 and 600 seconds")
        config.otp_timeout_seconds = args.timeout
    if args.download_dir:
        config.download_dir = args.download_dir.expanduser().resolve()
    if args.profile_dir:
        config.firefox_profile_dir = args.profile_dir.expanduser().resolve()
    if getattr(args, "excel_output_dir", None):
        config.excel_output_dir = args.excel_output_dir.expanduser().resolve()
    if getattr(args, "report_dir", None):
        config.excel_output_dir = args.report_dir.expanduser().resolve()
    if getattr(args, "xlsx_dir", None):
        config.download_dir = args.xlsx_dir.expanduser().resolve()
    if getattr(args, "image_dir", None):
        config.image_dir = args.image_dir.expanduser().resolve()
    if getattr(args, "report_dir", None):
        config.excel_output_dir = args.report_dir.expanduser().resolve()
    if getattr(args, "excel_template", None):
        config.excel_template_path = args.excel_template.expanduser().resolve()
    if getattr(args, "raw_xlsx", None):
        config.raw_xlsx_dir = args.raw_xlsx.expanduser().resolve().parent
    if getattr(args, "image", None):
        config.image_dir = args.image.expanduser().resolve().parent
    if getattr(args, "allow_connectivity_mutation", False):
        config.connectivity_allow_connect = True
        config.connectivity_allow_disconnect = True


def target_dates(
    query_date: date | None, start_date: date | None, end_date: date | None
) -> list[str]:
    """Resolve the notification date labels for a run."""
    if start_date and end_date:
        return [start_date.isoformat(), end_date.isoformat()]
    if query_date:
        return [query_date.isoformat()]
    return []


def notify(
    notifier: Notifier | None,
    event: str,
    *,
    mode: str,
    dates: list[str],
    started_at: float | None = None,
    record_count: int | None = None,
    error: BaseException | None = None,
) -> None:
    """Dispatch a lifecycle notification defensively.

    Notifications are best-effort: any failure here is warned about and
    swallowed so it can never change the pipeline result or exit code. Only
    short, allow-listed fields are sent; the output file path is intentionally
    excluded so messages stay concise.
    """
    if notifier is None:
        return
    try:
        elapsed = time.monotonic() - started_at if started_at is not None else None
        notifier.send_event(
            event,
            mode=mode,
            dates=dates,
            elapsed_seconds=elapsed,
            record_count=record_count,
            error_category=safe_error_category(error) if error is not None else None,
        )
    except Exception as exc:
        # Log the exception type only, matching the adapter's metadata-only
        # logging so a defensive-guard failure can never surface message text.
        logger.warning(
            "Lifecycle notification %s could not be dispatched: %s", event, type(exc).__name__
        )


async def main() -> int:
    """Main entry point."""
    args = parse_args()
    mode = getattr(args, "mode", "full")
    notifier: Notifier | None = None
    started_at: float | None = None
    notify_dates: list[str] = []

    try:
        # Load configuration
        config = load_config()
        apply_overrides(config, args)

        # Setup logging to file and console
        logs_dir = config.logs_dir or (config.excel_output_dir / "logs")
        setup_logging(logs_dir=logs_dir, log_level=args.log_level)

        # Interactive menu mode (like MRTG)
        if getattr(args, "menu", False):
            print("=" * 60)
            print("         TELKOMSEL CMP AUTOMATION")
            print("=" * 60)
            print("1. Scrape only (Export raw XLSX + capture dashboard image)")
            print("2. Generate only (Populate monthly Excel from raw/image)")
            print("3. Full pipeline (Scrape + Generate)")
            print("4. Exit")
            print("=" * 60)
            choice = input("Pilih menu (1-4): ").strip()
            if choice == "1":
                args.mode = "scrape"
            elif choice == "2":
                args.mode = "generate"
            elif choice == "3":
                args.mode = "full"
            elif choice == "4":
                return 0
            else:
                print("[!] Pilihan tidak valid.")
                return 1

        # Build the optional notifier before validation so configuration
        # errors can still be reported. Unconfigured notifications resolve to
        # None and stay a silent no-op; ``generate`` never notifies.
        if getattr(args, "mode", "full") in NOTIFIED_MODES:
            notifier = build_notifier(config)

        # Validate paths
        validate_paths(config)

        # Parse date arguments
        query_date = parse_date_arg(args.date)
        start_date = parse_date_arg(args.start_date)
        end_date = parse_date_arg(args.end_date)

        if start_date and end_date and start_date > end_date:
            raise ConfigurationError(
                f"--start-date ({start_date}) cannot be after --end-date ({end_date})"
            )

        mode = getattr(args, "mode", "full")
        notify_dates = target_dates(query_date, start_date, end_date)
        log_run_boundary("RUN START", f"mode={mode} date={query_date or start_date or 'today'}")

        # START is sent only after config/argument validation passed.
        if notifier is not None:
            started_at = time.monotonic()
            notify(notifier, EVENT_START, mode=mode, dates=notify_dates)

        # Run workflow
        result_path = await run_workflow(
            config,
            headed=args.headed,
            dry_run=args.dry_run,
            diagnose_export=args.diagnose_export,
            diagnose_auth=args.diagnose_auth,
            query_date=query_date,
            start_date=start_date,
            end_date=end_date,
            monitor_cycles=getattr(args, "monitor_cycles", None),
            connectivity_enabled=getattr(args, "enable_connectivity", False),
            allow_connectivity_mutation=getattr(args, "allow_connectivity_mutation", False),
            mode=mode,
            raw_xlsx=getattr(args, "raw_xlsx", None),
            image_path=getattr(args, "image", None),
        )

        logger.info("Success! Output: %s", result_path)
        log_run_boundary("RUN END", f"success output={result_path}")
        # The output path stays in the local logs only; notifications omit it.
        rec_count: int | None = None
        if result_path and result_path.exists():
            try:
                import openpyxl
                wb = openpyxl.load_workbook(result_path, read_only=True)
                # If monthly workbook with day sheets
                target_sheet = f"{(query_date or date.today()).day:02d}" if (query_date or date.today()) else None
                ws = wb[target_sheet] if target_sheet and target_sheet in wb.sheetnames else wb.active
                if ws is not None:
                    count = 0
                    for row in range(5, 55):
                        if ws.cell(row=row, column=4).value is not None:
                            count += 1
                        else:
                            break
                    rec_count = count if count > 0 else (ws.max_row - 1 if ws.max_row and ws.max_row > 1 else None)
                wb.close()
            except Exception:
                rec_count = None
        notify(
            notifier,
            EVENT_SUCCESS,
            mode=mode,
            dates=notify_dates,
            started_at=started_at,
            record_count=rec_count,
        )
        return 0

    except ConfigurationError as e:
        logger.error("Configuration error: %s", e)
        log_run_boundary("RUN END", f"failed error={e}")
        notify(notifier, EVENT_FAILED, mode=mode, dates=notify_dates, error=e)
        return 1
    except CMPAutomationError as e:
        logger.error("Automation error: %s", e)
        log_run_boundary("RUN END", f"failed error={e}")
        notify(
            notifier, EVENT_FAILED, mode=mode, dates=notify_dates, started_at=started_at, error=e
        )
        return 1
    except KeyboardInterrupt:
        # Interrupted runs intentionally send no terminal event; the run was
        # neither a success nor a normalized failure.
        logger.info("Interrupted by user")
        log_run_boundary("RUN END", "interrupted")
        return 130
    except Exception as e:
        logger.exception("Unexpected error")
        log_run_boundary("RUN END", "unexpected_error")
        notify(
            notifier, EVENT_FAILED, mode=mode, dates=notify_dates, started_at=started_at, error=e
        )
        return 1


def cli_main() -> None:
    """Synchronous entry point for setuptools."""
    try:
        code = asyncio.run(main())
    except KeyboardInterrupt:
        logger.info("Interrupted by user")
        code = 130
    sys.exit(code)


if __name__ == "__main__":
    cli_main()
