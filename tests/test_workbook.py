"""Comprehensive test suite for MonthlyWorkbookService and workbook utilities."""

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from PIL import Image

from cmp_automation.exceptions import WorkbookError
from cmp_automation.workbook import (
    clean_day_sheet,
    extract_data_from_xlsx,
    format_day_sheet_name,
    get_monthly_workbook_path,
    insert_dashboard_image,
    parse_numeric_value,
    parse_usage_bytes,
    read_lookup_table,
    write_lookup_table,
)


@pytest.fixture
def sample_image(tmp_path: Path) -> Path:
    """Create a temporary dummy PNG image."""
    img_path = tmp_path / "dashboard.png"
    img = Image.new("RGB", (100, 100), color="blue")
    img.save(img_path)
    return img_path


@pytest.fixture
def dummy_template(tmp_path: Path) -> Path:
    """Create a minimal mock template workbook with sheets '01'..'31' and 'Master'."""
    wb = Workbook()
    # openpyxl starts with 'Sheet'
    ws_master = wb.active
    ws_master.title = "Master"

    # Add lookup in Master
    ws_master["O3"] = "ICCID"
    ws_master["P3"] = "LOCATION"
    ws_master["O4"] = "8962000001"
    ws_master["P4"] = "HQ_MAIN"
    ws_master["O5"] = "8962000002"
    ws_master["P5"] = "BRANCH_NORTH"

    # Create day sheets '01'..'31'
    for day in range(1, 32):
        sname = f"{day:02d}"
        ws = wb.create_sheet(title=sname)
        # Add headers
        ws["B4"] = "NO"
        ws["C4"] = "DATE"
        ws["D4"] = "ICCID"
        ws["E4"] = "USAGE"
        ws["G4"] = "NO"
        ws["H4"] = "DATE"
        ws["I4"] = "ICCID"
        ws["J4"] = "LOCATION"
        ws["K4"] = "BYTES"
        ws["L4"] = "GB"

        # Lookup in day sheets
        ws["O3"] = "ICCID"
        ws["P3"] = "LOCATION"
        ws["O4"] = "8962000001"
        ws["P4"] = "HQ_MAIN"
        ws["O5"] = "8962000002"
        ws["P5"] = "BRANCH_NORTH"

    wb.create_sheet(title="MONTHLY")

    tpl_path = tmp_path / "template.xlsx"
    wb.save(tpl_path)
    wb.close()
    return tpl_path


def test_get_monthly_workbook_path(tmp_path: Path) -> None:
    """Test standardized monthly workbook path generation."""
    # Test with None (today)
    p_today = get_monthly_workbook_path(tmp_path)
    today = date.today()
    assert p_today.name == f"Daily-Data-Usage-M2M-{today.year:04d}{today.month:02d}.xlsx"

    # Test with specific date
    d = date(2026, 9, 15)
    p_custom = get_monthly_workbook_path(tmp_path, target_date=d)
    assert p_custom.name == "Daily-Data-Usage-M2M-202609.xlsx"

    # Test with datetime
    dt = datetime(2026, 12, 1, 10, 30)
    p_dt = get_monthly_workbook_path(tmp_path, target_date=dt, filename_prefix="CUSTOM-REPORT")
    assert p_dt.name == "CUSTOM-REPORT-202612.xlsx"


def test_format_day_sheet_name() -> None:
    """Test formatting day sheet names."""
    assert format_day_sheet_name(1) == "01"
    assert format_day_sheet_name(9) == "09"
    assert format_day_sheet_name(10) == "10"
    assert format_day_sheet_name(31) == "31"
    assert format_day_sheet_name(" 5 ") == "05"
    assert format_day_sheet_name(date(2026, 9, 7)) == "07"
    assert format_day_sheet_name(datetime(2026, 9, 25, 8, 0)) == "25"

    with pytest.raises(ValueError):
        format_day_sheet_name(0)
    with pytest.raises(ValueError):
        format_day_sheet_name(32)
    with pytest.raises(ValueError):
        format_day_sheet_name("abc")
    with pytest.raises(TypeError):
        format_day_sheet_name([1])  # type: ignore[arg-type]


def test_parse_numeric_value() -> None:
    """Test numeric string extraction."""
    assert parse_numeric_value(123) == 123.0
    assert parse_numeric_value(45.67) == 45.67
    assert parse_numeric_value("1,024.50 MB") == 1024.50
    assert parse_numeric_value("2,952,320,085") == 2952320085.0
    assert parse_numeric_value("-500.2") == -500.2
    assert parse_numeric_value(None) == 0.0
    assert parse_numeric_value("No numbers here") == 0.0


def test_parse_usage_bytes() -> None:
    """Test extracting usage bytes from dict or sequence."""
    assert parse_usage_bytes({"Total Data Usage (bytes)": 1000}) == 1000
    assert parse_usage_bytes({"total_data_usage_bytes": "2,500"}) == 2500
    assert parse_usage_bytes({"Usage": 5000}) == 5000
    assert parse_usage_bytes({"Byte": 777}) == 777
    assert parse_usage_bytes([1, "2026-09-01", "89620001", 12345]) == 12345
    assert parse_usage_bytes(["89620001", "3,456"]) == 3456
    assert parse_usage_bytes({}) == 0



def test_extract_data_from_xlsx(tmp_path: Path) -> None:
    """Test extracting records from raw CMP xlsx."""
    xlsx_path = tmp_path / "cmp_export.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["ICCID", "IMSI", "Total Data Usage (bytes)"])
    ws.append(["8962000001", "5100001", 1000000])
    ws.append(["8962000002", "5100002", 2000000])
    wb.save(xlsx_path)
    wb.close()

    records = extract_data_from_xlsx(xlsx_path)
    assert len(records) == 2
    assert records[0]["ICCID"] == "8962000001"
    assert records[0]["Total Data Usage (bytes)"] == 1000000

    # Non-existent file
    with pytest.raises(WorkbookError):
        extract_data_from_xlsx(tmp_path / "nonexistent.xlsx")

    # Invalid file
    bad_path = tmp_path / "corrupt.xlsx"
    bad_path.write_bytes(b"not an excel zip file")
    with pytest.raises(WorkbookError):
        extract_data_from_xlsx(bad_path)


def test_read_and_write_lookup_table() -> None:
    """Test reading and writing styled lookup table in O3:P140."""
    wb = Workbook()
    ws = wb.active

    lookup_data = {
        "8962000001": "HQ_MAIN",
        "8962000002": "BRANCH_NORTH",
    }
    write_lookup_table(ws, lookup_data)

    # Check header row 3
    assert ws["O3"].value == "ICCID"
    assert ws["P3"].value == "LOCATION"
    assert ws["O3"].font.bold is True

    # Check data rows
    assert ws["O4"].value == "8962000001"
    assert ws["P4"].value == "HQ_MAIN"
    assert ws["O5"].value == "8962000002"
    assert ws["P5"].value == "BRANCH_NORTH"

    # Read back
    read_back = read_lookup_table(ws)
    assert read_back == lookup_data
    wb.close()


def test_clean_day_sheet(sample_image: Path) -> None:
    """Test clean_day_sheet preserves layout and lookups while resetting data."""
    wb = Workbook()
    ws = wb.active

    # Populate some dummy data
    ws["B5"] = 999
    ws["C5"] = "2026-09-01"
    ws["D5"] = "8962000001"
    ws["E5"] = 5000000
    ws["O3"] = "ICCID"
    ws["P3"] = "LOCATION"
    ws["O4"] = "8962000001"
    ws["P4"] = "HQ"

    insert_dashboard_image(ws, sample_image)
    assert len(ws._images) == 1

    clean_day_sheet(ws, preserve_styles=True, preserve_lookup=True, clear_images=True)

    # Main table cleaned
    assert ws["B5"].value == 1
    assert ws["C5"].value is None
    assert ws["D5"].value is None
    assert ws["E5"].value is None

    # Total formulas set
    assert ws["D55"].value == "TOTAL"
    assert ws["E55"].value == "=SUM(E5:E54)"

    # Top 5 table cleaned
    assert ws["G5"].value == 1
    assert ws["H5"].value is None
    assert ws["L5"].value == "=K5/(1024^3)"

    # Summary table initialized
    assert ws["H12"].value == "TOTAL USAGE DAILY"
    assert ws["I12"].value == "Byte"
    assert ws["I13"].value == "=E55"
    assert ws["J13"].value == "=I13/(1024^3)"
    assert ws["K13"].value == "=COUNTA(D5:D54)"

    # Lookup preserved
    assert ws["O4"].value == "8962000001"
    assert ws["P4"].value == "HQ"

    # Images cleared
    assert len(ws._images) == 0
    wb.close()

